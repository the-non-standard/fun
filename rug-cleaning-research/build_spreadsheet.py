#!/usr/bin/env python3
"""Compile UK rug cleaning businesses into a formatted XLSX + CSV.

Data gathered via multi-source web search (Aug 2026). Directory sites and
company websites could not be fetched directly in this environment (network
egress policy blocked all outbound HTTPS), so every field below is what was
publicly surfaced through search result snippets and link URLs. Blank cells =
not found via search (usually still discoverable by visiting the site).
"""
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

COLS = [
    "Business Name", "Category", "Website", "Email", "Phone", "Address",
    "Facebook", "Instagram", "LinkedIn", "YouTube", "X/Twitter",
    "Reviews (as found)", "Coverage Area", "Source / Notes",
]

# Category key: SPEC = dedicated rug-cleaning specialist; C&R = carpet+rug cleaner;
# FRAN = national franchise/platform.
ROWS = [
    # ---------- LONDON & SOUTH EAST ----------
    ["James Barclay", "SPEC", "https://jamesbarclay.co.uk", "info@jamesbarclay.co.uk", "0203 174 2427",
     "Oliver Business Park, Park Royal, London NW10 7JB", "", "", "", "", "",
     "Google reviews page; website testimonials", "National (free UK collection)",
     "Oriental/Persian specialist; Guild of Master Craftsmen"],
    ["Oriental Rug Services Ltd", "SPEC", "https://orientalrugservices.com", "info@orientalrugservices.com",
     "020 7625 6931 / 07875 694100", "Kilburn Square, London NW6 6PP",
     "https://www.facebook.com/OrientalRugServicesKilburn/", "https://www.instagram.com/orientalrugservicesltd/",
     "", "", "", "", "Greater London", "Specialist rug cleaning & repair"],
    ["The Persian Rug Cleaning Company", "SPEC", "https://persian-rug-cleaning.london", "", "",
     "South Kensington, London", "", "", "", "", "", "", "London",
     "Family-run Persian/Oriental specialist, 25+ yrs"],
    ["RugMaster", "SPEC", "https://rugmaster.uk", "", "0208 341 9191", "Barnet, London",
     "https://www.facebook.com/rugsmaster/", "", "", "", "", "Yell reviews (positive)", "North London",
     "Family-owned 3 generations; clean/repair/valuation"],
    ["London Persian Rug Company", "SPEC", "https://londonpersianrugcompany.co.uk", "", "", "London",
     "", "", "", "", "", "", "London", "Traditional hand-cleaning/repair, 20+ yrs"],
    ["Oriental Rug Cleaning Company", "SPEC", "https://orientalrugcleaningcompany.co.uk", "", "", "London",
     "", "", "", "", "", "", "London", "Oriental rug specialist"],
    ["Orientalist Rug", "SPEC", "https://orientalrugservicesltd.co.uk", "", "",
     "London", "https://www.facebook.com/OrientalistRug/", "", "", "", "", "", "London", "Web search"],
    ["The Oriental Rug Repair Company (London EC2)", "SPEC", "", "", "", "London EC2",
     "", "", "", "", "", "", "London", "Guild of Master Craftsmen listing"],
    ["ProLux Cleaning", "C&R", "https://www.proluxcleaning.co.uk/rug-cleaning", "", "", "London",
     "", "", "", "", "", "", "London", "15+ yrs; Persian/handmade/Oriental"],
    ["Anyclean", "C&R", "https://www.anyclean.co.uk", "", "", "London", "", "", "", "", "", "",
     "London", "NCCA member"],
    ["Rug Cleaning London", "C&R", "https://www.rugcleaninglondon.org.uk", "", "", "London",
     "", "", "", "", "", "", "London", "Web search"],
    ["Absolutely Fabulous Persian & Oriental Rug Cleaning", "SPEC", "", "", "", "London & surrounding",
     "", "", "", "", "", "", "London", "Silk/wool/antique/Oriental/Persian"],
    ["AAAClean", "C&R", "https://www.aaaclean.co.uk", "", "", "Kent", "", "", "", "", "", "",
     "Kent / London / East Sussex", "In-situ + commercial"],
    ["Carpet Clean Kent", "C&R", "https://carpetcleankent.co.uk", "", "", "Kent",
     "", "", "", "", "", "", "Kent / Sussex / Surrey", "20+ yrs"],
    ["Apple Clean", "C&R", "https://www.appleclean.co.uk", "", "", "Sussex", "", "", "", "", "", "",
     "Sussex / Surrey / Hampshire / Kent", "35+ yrs"],
    ["My Carpet Doctor", "C&R", "https://my-carpet-doctor.webnode.co.uk", "", "", "South East",
     "", "", "", "", "", "", "London / Kent / Surrey / Essex / Sussex", "Web search"],
    ["Oriental Rug Spa", "SPEC", "https://www.orientalrugspa.co.uk", "", "01276 423 150",
     "Camberley, Surrey", "", "", "", "", "", "Website testimonials", "National (collection: Newcastle, Birmingham, Devon, etc.)",
     "Dedicated facility in Camberley; Persian/Oriental/Turkish + repairs"],

    # ---------- EAST OF ENGLAND ----------
    ["Suffolk & Essex Rug Care", "SPEC", "https://www.suffolkandessexrugcare.co.uk", "enquiries@suffolkrugcare.co.uk",
     "01473 396036 / 07966 586852", "Unit E, Sycamore Farm, Somersham Road, Ipswich IP8 4NN",
     "https://www.facebook.com/suffolkrugcare/", "", "", "", "", "Facebook (4); Checkatrade testimonials",
     "Suffolk / Essex / Norfolk / London", "Oriental/Persian/handmade/machine-made"],
    ["Oriental Rugs of Norwich Ltd", "SPEC", "https://orientalrugsofnorwich.co.uk", "ornltd@hotmail.co.uk",
     "01603 633520", "4 Bedford Street, Norwich NR2 1AR", "https://www.facebook.com/OrientalRugsofNorwich/",
     "", "", "", "", "", "Norfolk / Suffolk / Cambridgeshire", "35+ yrs, own workshop; Companies House 04563930"],
    ["Rugs 2 Riches", "C&R", "https://www.rugs2riches.com", "", "", "Norwich, Norfolk",
     "", "", "", "", "", "", "Norwich", "Carpet/rug/upholstery"],
    ["Carpet Clean Norwich", "C&R", "https://www.carpetcleannorwich.co.uk", "", "", "Norwich",
     "", "", "", "", "", "", "Norwich / Norfolk", "Web search"],
    ["Clean Living Norfolk", "C&R", "https://www.cleanlivingnorfolk.co.uk", "", "", "Norfolk",
     "", "", "", "", "", "", "Norfolk / Suffolk", "Web search"],

    # ---------- SOUTH / SOUTH WEST ----------
    ["Rug Cleaning Works", "SPEC", "https://www.rugcleaningworks.co.uk", "mail@rugcleaningworks.co.uk",
     "01453 836400", "Unit 11, Nailsworth Mills Estate, Avening Road, Nailsworth, Stroud GL6 0BS",
     "", "", "", "", "", "", "Cotswolds (Bath, Bristol, Cheltenham, Gloucester, Worcester, Swindon)",
     "Independent family firm, est 2002"],
    ["Captain Rug Wash", "SPEC", "https://captainrugwash.co.uk", "andi@plymouthrugcleaning.co.uk",
     "01752 406210", "1 Hogarth Walk, Plymouth PL9 8ET", "https://www.facebook.com/captainrugwash/",
     "", "", "", "", "Facebook (1,875 likes); Yell (6)", "Plymouth / Devon / Cornwall / South West",
     "Founded 1988; Andi Hill, IICRC + WoolSafe"],
    ["Masterclean Devon", "C&R", "https://www.mastercleanexeter.uk", "", "", "Exeter, Devon",
     "", "", "", "", "", "", "Devon", "HWE + Oriental rug cleaning"],
    ["Rug Nurse", "SPEC", "https://rugnurse.co.uk", "inforugnurse@gmail.com", "07900 867497",
     "Oxford (serves Devon by collection)", "https://www.facebook.com/TheRugNurse/",
     "https://www.instagram.com/rugnurse/", "", "", "", "Facebook (3)", "Devon / Oxford", "Persian rug cleaning & repairs"],
    ["Majestic Cleaning SW", "C&R", "https://majesticcleaningsw.co.uk", "", "", "Devon / Cornwall",
     "", "", "", "", "", "", "Devon / Cornwall", "Web search"],
    ["Devon Cleaning", "C&R", "https://www.devoncleaning.com", "", "", "Plymouth / Exeter, Devon",
     "", "", "", "", "", "", "Devon / Cornwall", "Window/carpet/contract"],
    ["Cotswold Carpet Cleaners", "C&R", "https://cotswoldcarpetcleaners.co.uk", "", "", "Cotswolds",
     "", "", "", "", "", "", "Cotswolds", "Rug cleaning service"],
    ["Rug Repair Co (Oxfordshire & Cotswolds)", "SPEC", "https://cotswolds.rugrepairco.co.uk", "", "",
     "Oxfordshire / Cotswolds", "", "", "", "", "", "", "Oxfordshire / Cotswolds", "Oriental rug cleaning & repair"],
    ["Rug Clean Southampton", "SPEC", "http://rug-clean.co.uk", "", "0800 695 6956 / 023 8001 6671",
     "2 Northlands Gdns, Southampton SO15 2NL", "", "", "", "", "", "", "Southampton / Hampshire",
     "Hand-wash Oriental/Persian; clean/repair/dye/restore"],
    ["Hampshire Carpet Care", "C&R", "https://www.hampshirecarpetcare.co.uk", "", "", "Southampton",
     "", "", "", "", "", "", "Southampton / Hampshire", "WoolSafe-approved on-site"],
    ["Advanced Carpet Cleaning", "C&R", "https://www.advanced-carpet-cleaning.co.uk", "", "", "Southampton",
     "", "", "", "", "", "", "Southampton", "Web search"],
    ["Mountain Rug Cleaning", "SPEC", "https://mountainrugcleaning.co.uk", "info@mountainrugcleaning.co.uk",
     "", "Hampshire (Ltd reg: 5 High St, Westbury On Trym, Bristol BS9 3BY)",
     "https://www.facebook.com/mountainrugclean/", "", "", "https://www.youtube.com/@MountainRugCleaning",
     "https://twitter.com/mountrugclean", "", "Hampshire / South West", "Water-immersion rug spa; pickup & delivery"],
    ["Cleaning Rugs (Bournemouth)", "SPEC", "http://cleaning-rugs.co.uk", "", "01202 906135",
     "Bournemouth", "", "", "", "", "", "", "Bournemouth / Dorset", "Rug clean/repair/restore"],
    ["Elite Carpet & Cleaning Services", "C&R", "https://elitecarpetandcleaningservices.co.uk", "", "",
     "Bournemouth / Poole / Ferndown", "", "", "", "", "", "", "Dorset", "15+ yrs"],

    # ---------- MIDLANDS ----------
    ["Jacoby Rugs", "SPEC", "https://jacobyrugs.co.uk", "rugs@jacobyrugs.co.uk", "01788 247077",
     "Workshops near Warwick (Warwickshire) & near Edinburgh", "", "", "", "", "", "Website testimonials",
     "Midlands / Birmingham / Warwickshire / Scotland", "Family-run Oriental cleaning & repair, two workshops"],
    ["Cleaning Direct", "SPEC", "", "", "", "Edgbaston, Birmingham", "", "", "", "", "", "",
     "Birmingham", "25 yrs professional rug cleaners (shop in Edgbaston)"],
    ["Upholstery Carpet Cleaners", "C&R", "https://www.upholsterycarpetcleaners.co.uk", "", "", "Birmingham",
     "", "", "", "", "", "", "Birmingham / West Midlands", "Persian/wool/Oriental/silk"],
    ["The Carpet Cleaner Birmingham", "C&R", "https://thecarpetcleanerbirmingham.co.uk", "", "", "Birmingham",
     "", "", "", "", "", "", "Birmingham", "Web search"],
    ["S.P. Carpet & Upholstery Care (Nottingham Rug Cleaners)", "SPEC", "https://www.nottinghamrugcleaners.co.uk",
     "forabetterclean@aol.com", "0115 971 8323", "Nottingham", "", "", "", "", "", "",
     "Nottingham / Derby", "30 yrs rug/carpet expertise"],
    ["Flawless Group", "C&R", "https://www.flawlessltd.co.uk", "", "0800 002 5959", "Nottingham",
     "", "", "", "", "", "", "Nottingham", "Rug cleaners"],
    ["Clean Carpets Nottingham", "C&R", "https://www.cleancarpetsnottingham.co.uk", "", "07852 328686",
     "Nottingham", "", "", "", "", "", "", "Nottingham", "Web search"],
    ["Just Like Nu (Nottingham Carpet Cleaning)", "C&R", "https://www.nottinghamcarpetcleaning.com", "",
     "07760 281510", "Nottingham", "", "", "", "", "", "", "Nottinghamshire", "Web search"],
    ["CFFC Services", "C&R", "https://www.cffcservices.co.uk", "", "", "Nottingham / Derby",
     "", "", "", "", "", "", "Nottinghamshire / Derbyshire", "Carpet cleaners"],
    ["Pristine Clean", "C&R", "https://www.pristine-carpets.co.uk", "", "", "Leicester / Derby / Nottingham",
     "", "", "", "", "", "", "East Midlands", "Web search"],
    ["Neat Carpet Cleaning", "C&R", "https://neatcarpetcleaning.co.uk", "", "", "Leicester / Nottingham",
     "", "", "", "", "", "", "East Midlands", "5-star family run"],
    ["Bailey's Floor Care (Floor Care Specialists)", "C&R", "https://www.floorcarespecialists.co.uk", "",
     "0115 775 0191", "Derby", "", "", "", "", "", "", "Derby / Leicester / Northampton / Nottingham", "24 yrs"],

    # ---------- NORTH WEST ----------
    ["Arcadia Rug Spa", "SPEC", "https://arcadiarugspa.co.uk", "info@arcadiarugspa.co.uk",
     "01606 882712 / 07400 159528", "Unit 7, Hartford Business Centre, Chester Road, Hartford, Northwich CW8 2AB",
     "https://www.facebook.com/Arcadiarugspa/", "https://www.instagram.com/arcadiarugspa/", "", "", "",
     "Trustpilot 4.0 (32); Facebook (26)", "Cheshire / Manchester / Liverpool / Warrington",
     "Independent family run; off-site workshop"],
    ["Cheshire Rug Cleaning", "SPEC", "https://cheshirerugcleaning.co.uk", "info@cheshirerugcleaning.co.uk",
     "01606 302500", "Unit 6, Atlantic Trading Park, School Road, Meadowbank, Winsford CW7 2PG",
     "https://www.facebook.com/1720149511579938", "", "", "", "", "5.0 (62); Facebook (34)", "Cheshire",
     "Specialist workshop"],
    ["Rug Cleaning Experts", "SPEC", "https://www.rugcleaningexperts.co.uk", "", "", "Bury, Greater Manchester",
     "https://www.facebook.com/RugCleaningExperts/", "", "", "", "", "Self-reported 5 stars, 500+ customers",
     "Manchester / North West", "Family-run; off-site Persian rug washing"],
    ["Mike Bradburn Oriental Rug Cleaning (RugRepair.org)", "SPEC", "https://rugrepair.org", "", "",
     "North West England (Manchester)", "", "", "", "", "", "", "Manchester / Cheshire / Lancashire / Merseyside / Cumbria / Derbyshire",
     "Hand-clean specialist; Iranian repair team"],
    ["Wrennalls Group", "C&R", "https://wrennalls.com", "", "", "Preston, Lancashire",
     "", "", "", "", "", "", "Chester / Bolton / Preston / Bury", "30+ yrs floor care"],
    ["Stain Kings", "C&R", "https://stainkings.com", "", "", "Manchester", "", "", "", "", "", "",
     "Manchester", "Persian/Oriental rug cleaning"],
    ["Manchester House Cleaning Services", "C&R", "https://www.manchesterhousecleaningservices.co.uk", "",
     "0161 676 2888", "Manchester", "", "", "", "", "", "", "Manchester", "Web search"],
    ["Absolute Clean Manchester", "C&R", "https://absolutecleanmanchester.co.uk", "", "", "Manchester",
     "", "", "", "", "", "", "Manchester / Oldham / Rochdale", "Collection & delivery"],

    # ---------- YORKSHIRE & NORTH EAST ----------
    ["Star Fabric Care", "C&R", "https://starfabriccare.co.uk", "", "", "Leeds", "", "", "", "", "", "",
     "Leeds / Bradford / Harrogate / York", "Family-run since 1999"],
    ["Hoyles Cleaning Specialists", "C&R", "https://hoylescleaningspecialists.co.uk", "", "", "Yorkshire",
     "", "", "", "", "", "", "Bradford / Halifax / Huddersfield / Harrogate / Skipton", "Rug cleaning specialist"],
    ["Esteamed Carpet Cleaning", "C&R", "https://esteamedcarpetcleaning.co.uk", "", "", "West Yorkshire",
     "", "", "", "", "", "", "Leeds / Bradford / Ilkley / Harrogate", "Web search"],
    ["Trust Cleaner", "C&R", "https://trustcleaner.co.uk", "", "", "Yorkshire", "", "", "", "", "", "",
     "Leeds / Sheffield / York / Harrogate", "Yorkshire cleaning specialists"],
    ["Chem-Dry Harrogate York & Leeds", "FRAN", "https://www.chemdryharrogateyorkleeds.co.uk", "", "",
     "Roecliffe, Boroughbridge", "", "", "", "", "", "", "Harrogate / York / Leeds", "Chem-Dry franchise; rug facility"],
    ["Oriental Rug Cleaning (York/Leeds - rugrepair.org)", "SPEC", "https://york.rugrepair.org", "", "",
     "Leeds / York", "", "", "", "", "", "", "Leeds / York / Sheffield / Doncaster / Harrogate / Hull",
     "Hand cleaning, traditional methods"],
    ["Oriental Rug Spa (Newcastle service)", "SPEC", "https://www.orientalrugspa.co.uk/rug-cleaning-newcastle-2/",
     "", "", "Serves Newcastle (HQ Camberley)", "", "", "", "", "", "", "Newcastle / North East",
     "Collection-based; see Oriental Rug Spa"],
    ["ServiceMaster Clean (Newcastle & Gateshead)", "FRAN", "https://www.servicemasterclean.co.uk", "", "",
     "Newcastle / Gateshead", "", "", "", "", "", "", "North East", "Franchise; rug spa"],

    # ---------- SCOTLAND ----------
    ["The Oriental Rug Repair Co (ORRC)", "SPEC", "https://www.orrc.co.uk", "enquiries@orrc.co.uk",
     "0131 221 6527", "28 Rutland Square, Edinburgh EH1 2BW", "https://www.facebook.com/TheOrientalRugRepairCo/",
     "", "", "", "", "Positive across platforms", "Edinburgh / Aberdeen / Glasgow / Scotland",
     "Family business; Guild of Master Craftsmen"],
    ["Rug Cleaner Scotland", "SPEC", "https://www.rugcleanerscotland.co.uk", "david@rugcleanerscotland.co.uk",
     "07500 149433", "Unit 4, 3 Cloberfield, Glasgow G62 7LN", "https://www.facebook.com/rugcleanerscotland/",
     "", "", "", "", "Facebook (1,217 likes); Yell (7)", "Scotland (Glasgow, Edinburgh)",
     "Family-owned dedicated rug cleaners, 16+ yrs"],
    ["Kleen-Dri", "C&R", "https://kleendri.co.uk", "hello@kleendri.co.uk", "07886 856024", "Glasgow",
     "", "", "", "", "", "", "Glasgow / Lanarkshire", "Rug cleaning Glasgow"],
    ["Pro Carpet Cleaning", "C&R", "https://pro-carpetcleaning.co.uk", "", "", "Edinburgh / Fife / Dundee / Tayside",
     "", "", "", "", "", "", "East Scotland", "Rug restoration"],
    ["NOVA CLEAN", "C&R", "https://www.cleaning-company-aberdeen.com", "", "", "Aberdeen / Aberdeenshire",
     "", "", "", "", "", "", "Aberdeen", "Home call-outs + commercial"],

    # ---------- WALES ----------
    ["CSB Cleaning", "C&R", "https://csbcleaning.co.uk", "", "01656 859611", "Bridgend / Cardiff / South Wales",
     "https://www.facebook.com/Cleaning.Solutions.Bridgend/", "", "", "", "", "Facebook (1,867 likes)",
     "South Wales (Cardiff, Swansea, Bridgend)", "Award-winning family run (Kevin & Mair)"],
    ["Edwards Jeffery", "C&R", "https://edwardsjefferycarpetcleaning.co.uk", "", "", "Cardiff / South Wales",
     "", "", "", "", "", "", "South Wales", "Multi-award-winning, fully insured"],
    ["Ultra Clean Services", "C&R", "https://www.ultracleanservices.co.uk", "", "",
     "Cardiff / Newport / Caerphilly", "", "", "", "", "", "", "South Wales", "Professional rug cleaning"],
    ["Mr Jones' Rug Cleaning Spa", "SPEC", "", "", "", "185a Cardiff Rd, Newport", "", "", "", "", "", "",
     "Newport / South Wales", "Persian/Oriental to machine-made; dye bleed/stain/odour"],
    ["Cardiff Carpet Cleaning Co", "C&R", "", "", "", "Cardiff, South Wales", "", "", "", "", "", "",
     "South Wales / M4 corridor", "IICRC trained (via Yell)"],
    ["Clean Rescue", "C&R", "https://cleanrescue.uk", "", "", "Wales", "", "", "", "", "", "", "Wales", "Web search"],

    # ---------- NORTHERN IRELAND ----------
    ["The Rug Spa (Belfast)", "SPEC", "https://www.therugspa.co.uk", "", "02890 814063 / 0800 988 3362",
     "27 Oakwood Ave, Carryduff, Belfast BT8 8SW", "", "", "", "", "", "", "Belfast / Northern Ireland",
     "10+ yrs; domestic & commercial"],
    ["Elite Carpet Cleaners (Belfast)", "C&R", "https://elitecarpetcleaners.com", "", "", "Belfast",
     "", "", "", "", "", "", "Belfast / Newtownabbey / Carrickfergus", "Traditional/Oriental/Persian/kilim"],
    ["Ultra Clean Belfast", "C&R", "https://www.ultracleanbelfast.co.uk", "", "07974 827737", "Belfast",
     "", "", "", "", "", "", "Belfast / North Down", "25+ yrs"],
    ["Chem-Dry Northern", "FRAN", "https://www.chemdrynorthern.co.uk", "", "", "Belfast",
     "", "", "", "", "", "", "Belfast / Northern Ireland", "Chem-Dry franchise"],

    # ---------- NATIONAL / MULTI-REGION ----------
    ["The Rug Laundry", "SPEC", "https://www.theruglaundry.co.uk", "hello@theruglaundry.co.uk",
     "01269 842497 / 0800 002 9543", "Unit 16, Capel Hendre Industrial Estate, Ammanford SA18 3SJ",
     "https://www.facebook.com/theruglaundry.co.uk/", "https://www.instagram.com/theruglaundry/", "", "", "",
     "Trustpilot (26); Facebook 98% (43)", "National (collection network + drop-off points)",
     "Family-run 20+ yrs; WoolSafe member"],
    ["Oriental Rug Repair Co", "SPEC", "https://orientalrugrepair.co.uk", "", "", "UK",
     "", "", "", "", "", "", "UK", "Guild of Master Craftsmen"],
    ["Rug Restoration", "SPEC", "https://rugrestoration.co.uk", "", "", "UK", "", "", "", "", "", "",
     "UK", "Restoration / cleaning / repair"],
    ["eMop", "FRAN", "https://www.emop.co.uk", "", "", "London / national", "", "", "", "", "", "",
     "National (Birmingham, Manchester, Leeds, ...)", "On-demand cleaning platform"],
    ["Hello Services", "FRAN", "https://helloservices.co.uk", "", "", "National", "", "", "", "", "", "",
     "National (Liverpool, Bristol, Sheffield, Nottingham, Newcastle)", "Booking platform"],
    ["Cleaning Doctor", "FRAN", "https://cleaningdoctor.net", "", "", "National (franchise)", "", "", "", "", "", "",
     "National franchise (Birmingham, Edinburgh, Belfast, Norwich, ...)", "Franchise network"],

    # ---------- BRISTOL / SOMERSET (social-first listings) ----------
    ["Bristol Rug Cleaning", "SPEC", "", "", "", "Clevedon / Bristol",
     "https://www.facebook.com/bristolrugcleaning/", "", "", "", "", "", "Bristol area", "Facebook business page"],
    ["Rug Cleaning Weston-super-Mare", "SPEC", "", "", "", "Weston-super-Mare / Bristol",
     "https://www.facebook.com/rugcleaningWSM/", "", "", "", "", "", "Bristol / Somerset", "Facebook business page"],
    ["My Care Carpet Cleaners", "C&R", "https://www.mycarecarpetcleaners.co.uk", "", "", "Bristol",
     "", "", "", "", "", "", "Bristol", "Web search"],
]

# ---- Write CSV ----
with open("uk_rug_cleaning_businesses.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(COLS)
    w.writerows(ROWS)

# ---- Write XLSX ----
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "UK Rug Cleaning Businesses"

HEAD_FILL = PatternFill("solid", fgColor="1F3B4D")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
SPEC_FILL = PatternFill("solid", fgColor="E8F3EC")
FRAN_FILL = PatternFill("solid", fgColor="FFF4E5")
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(vertical="top", wrap_text=True)

ws.append(COLS)
for c in ws[1]:
    c.fill = HEAD_FILL
    c.font = HEAD_FONT
    c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    c.border = BORDER

for r in ROWS:
    ws.append(r)

link_cols = {COLS.index(h) + 1 for h in ("Website", "Facebook", "Instagram", "LinkedIn", "YouTube", "X/Twitter")}
for ridx in range(2, ws.max_row + 1):
    cat = ws.cell(row=ridx, column=2).value
    fill = SPEC_FILL if cat == "SPEC" else (FRAN_FILL if cat == "FRAN" else None)
    for cidx in range(1, len(COLS) + 1):
        cell = ws.cell(row=ridx, column=cidx)
        cell.alignment = WRAP
        cell.border = BORDER
        if fill:
            cell.fill = fill
        if cidx in link_cols and isinstance(cell.value, str) and cell.value.startswith("http"):
            cell.hyperlink = cell.value
            cell.font = Font(color="0563C1", underline="single", size=10)
        else:
            cell.font = Font(size=10)

widths = [30, 9, 34, 30, 22, 40, 42, 40, 16, 34, 26, 26, 40, 46]
for i, wdt in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = wdt

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws.max_row}"

# ---- README / method sheet ----
ws2 = wb.create_sheet("Method & Notes")
notes = [
    ["UK Rug Cleaning Businesses — research notes"],
    [""],
    [f"Total businesses listed: {len(ROWS)}"],
    ["Compiled: August 2026, via multi-source web search."],
    [""],
    ["Category key:"],
    ["  SPEC = dedicated rug-cleaning specialist (green rows)"],
    ["  C&R  = carpet & upholstery cleaner that also cleans rugs"],
    ["  FRAN = national franchise / booking platform (amber rows)"],
    [""],
    ["IMPORTANT — data completeness:"],
    ["  Individual company websites and directory sites (Yell, Checkatrade, Bark,"],
    ["  Trustpilot, FreeIndex, Cylex) could NOT be opened directly in this"],
    ["  environment — the network policy blocked all outbound page fetching."],
    ["  Every field here is what surfaced through search-result snippets and link"],
    ["  URLs. Blank cells mean 'not found via search', not 'does not exist' —"],
    ["  emails and Instagram/LinkedIn in particular are usually still on the site."],
    [""],
    ["To make this exhaustive (every UK business), the reliable routes are:"],
    ["  1. Google Places / Maps API — query 'rug cleaning' per town/postcode."],
    ["  2. Yell.com, Checkatrade, Bark, TrustATrader, Thomson Local, FreeIndex,"],
    ["     Cylex, Yelp UK — scrape/export their rug-cleaning category listings."],
    ["  3. NCCA (ncca.co.uk) & WoolSafe (woolsafe.org) member directories —"],
    ["     the trade bodies for this niche."],
    ["  4. Companies House (find-and-update.company-information.service.gov.uk)"],
    ["     SIC code 96010 for registered entities."],
    [""],
    ["Directories seen referencing these businesses (good next-step sources):"],
    ["  ncca.co.uk/trustedlocalcleaners, woolsafe.org, yell.com, trustatrader.com,"],
    ["  checkatrade.com, freeindex.co.uk, cylex-uk.co.uk, uk.trustpilot.com,"],
    ["  findacraftsman.com (Guild of Master Craftsmen), listedin.co.uk"],
]
for row in notes:
    ws2.append(row)
ws2["A1"].font = Font(bold=True, size=14)
ws2["A3"].font = Font(bold=True)
ws2["A11"].font = Font(bold=True, color="B00000")
ws2.column_dimensions["A"].width = 78

wb.save("uk_rug_cleaning_businesses.xlsx")
print(f"Wrote {len(ROWS)} businesses to XLSX + CSV")
